# ----------------------------------------------------------------------
# views_asistente.py — endpoints del asistente AI.
# ----------------------------------------------------------------------
# Flujo principal (POST /api/asistente/mensaje/):
#   1. Recibimos mensaje del user. Lo guardamos en BD.
#   2. Construimos contexto: últimos N mensajes (cap configurable).
#   3. Llamamos al LLM con tools disponibles.
#   4. Si pide tool_use → ejecutamos tool (con permisos del user),
#      guardamos resultado en BD, volvemos al paso 3 (loop hasta tope).
#   5. Cuando el modelo responde texto final, lo guardamos y devolvemos.
#
# Persistencia: TODOS los mensajes se guardan en BD. El contexto enviado
# al LLM se limita a `AsistenteConfig.contexto_max_mensajes` (default 20).
# ----------------------------------------------------------------------

import json
import logging
from typing import Any

from django.contrib.auth.decorators import login_required
from django.db.models import Q
from django.http import JsonResponse
from django.views.decorators.http import require_http_methods

from .models import (
    AsistenteConfig, ConversacionAsistente, MensajeAsistente,
)
from .asistente_provider import chat, AsistenteError
from .asistente_tools import execute_tool, tools_for_user

log = logging.getLogger(__name__)

# Tope de iteraciones tool-use → respuesta. Evita loops infinitos.
MAX_TOOL_ITERATIONS = 6


def _get_or_create_conv(user):
    conv, _ = ConversacionAsistente.objects.get_or_create(usuario=user)
    return conv


def _msg_to_dict(m: MensajeAsistente) -> dict:
    """Convierte un MensajeAsistente a formato OpenAI-compatible para mandar al LLM."""
    if m.role == 'tool':
        return {
            'role': 'tool',
            'tool_call_id': m.tool_call_id,
            'content': m.contenido or '',
        }
    if m.role == 'assistant' and m.tool_name:
        # Llamada a tool del assistant. Reconstruir el tool_calls array.
        return {
            'role': 'assistant',
            'content': m.contenido or None,
            'tool_calls': [{
                'id': m.tool_call_id,
                'type': 'function',
                'function': {
                    'name': m.tool_name,
                    'arguments': m.tool_args_json or '{}',
                },
            }],
        }
    return {'role': m.role, 'content': m.contenido or ''}


def _build_context(conv: ConversacionAsistente, config: AsistenteConfig) -> list[dict]:
    """Toma los últimos N mensajes y los formatea para el LLM."""
    n = config.contexto_max_mensajes or 20
    qs = conv.mensajes.order_by('-fecha')[:n]
    # Revertir para orden cronológico
    msgs = list(reversed(list(qs)))
    return [_msg_to_dict(m) for m in msgs]


def _user_context_block(user) -> str:
    """Datos relevantes del user que el asistente DEBE saber de entrada.
    No son secretos — son cosas que mejoran las respuestas y dan personalidad
    (saludar por nombre, conocer rol, saber su carga actual).
    """
    from .models import TodoItem, UserProfile
    from django.utils import timezone

    full = user.get_full_name() or user.username
    first = (user.first_name or user.username).strip()
    last = (user.last_name or '').strip()
    # Rol
    rol_display = 'vendedor'
    try:
        prof = UserProfile.objects.filter(user=user).first()
        if prof and prof.rol:
            rol_display = prof.get_rol_display() if hasattr(prof, 'get_rol_display') else str(prof.rol)
    except Exception:
        pass
    # Oportunidades activas del user
    try:
        opp_activas = TodoItem.objects.filter(usuario=user).exclude(
            Q(etapa_completa__icontains='perdido') |
            Q(etapa_completa__icontains='cancelad')
        ).count()
    except Exception:
        opp_activas = '?'
    # Fecha actual
    now = timezone.now()
    meses_es = ['', 'enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio',
                'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']
    fecha_str = f'{meses_es[now.month]} {now.year}'

    return (
        f'\n\n## Contexto del usuario actual\n'
        f'- Nombre: **{full}** (puedes llamarle "{first}" en conversación informal)\n'
        f'- Usuario en el sistema: @{user.username}\n'
        f'- Rol: {rol_display}\n'
        f'- Oportunidades activas a su nombre: {opp_activas}\n'
        f'- Fecha de hoy: {fecha_str}\n'
    )


def _system_prompt(user, config: AsistenteConfig) -> dict:
    """System prompt — personalidad + reglas + contexto del user."""
    custom = (config.system_prompt or '').strip()
    if custom:
        # Si el admin definió un prompt custom, lo respetamos pero le agregamos
        # el contexto del user al final.
        return {'role': 'system', 'content': custom + _user_context_block(user)}

    base = (
        f'Eres {config.nombre}, asistente AI del CRM de IAMET — una empresa '
        'mexicana de soluciones de tecnología, automatización industrial y '
        'sistemas de identificación. Tu rol es **consultor de negocio** '
        'para vendedores y supervisores: reportes claros, proyecciones, '
        'evaluaciones de desempeño y análisis del pipeline. Hablas español '
        'mexicano natural, tono ejecutivo pero cercano. Humor sutil cuando '
        'cae bien — nunca payaso. Eres conciso por default; expandes solo '
        'cuando el dato lo amerita.\n\n'

        '## Reglas duras (NO negociables)\n'
        '1. **NUNCA inventes datos.** Si necesitas información del CRM, USA '
        'las herramientas. Si una tool devuelve lista vacía o error, dilo '
        'claro sin disfrazar.\n'
        '2. **Cuando listes oportunidades**, SIEMPRE formatea cada una como '
        'link clickable usando el formato custom `[Título — Cliente — '
        '$Monto](opp:ID)`. El sistema convierte esos links en botones que '
        'abren la oportunidad en el CRM. Ejemplo:\n'
        '   `- [Renovación Carl Zeiss — Carl Zeiss México — $1,250,000](opp:1198)`\n'
        '3. **Para reportes con varias columnas** (vendedor + cliente + '
        'monto + etapa, etc.), usa **tablas markdown** con encabezados:\n'
        '   ```\n'
        '   | Vendedor | Opp creadas | Monto ganado |\n'
        '   |----------|-------------|--------------|\n'
        '   | Ana López | 12 | $1,450,000 |\n'
        '   ```\n'
        '4. **Formato de montos**: siempre con $, separadores de miles y '
        'sufijo MXN cuando es relevante. Ej. `$1,250,000 MXN`.\n'
        '5. **No expongas IDs internos** crudos como texto suelto (ej. '
        '"#12345"). Los IDs solo van dentro de `(opp:ID)` para los links.\n'
        '6. **Headings con `###`** SOLO para secciones de respuestas largas. '
        'Respuestas cortas no necesitan headings. NO uses `####` ni más #, '
        'usa máximo `###`.\n'
        '6b. **NUNCA inventes URLs.** Si una tool devuelve un link, cópialo '
        'literal. No agregues parámetros ni cambies dominios.\n'
        '6c. **Forecast del mes**: cuando el user diga "cómo cerraremos '
        'el mes" SIN especificar mes, llama `forecast_cierre` con los '
        'defaults (toma mes anterior + mes actual automáticamente — las '
        'opp que ya debían cerrar pero siguen abiertas + las que cierran '
        'este mes). NO le pases todo el pipeline.\n'
        '6d. **Lista de oportunidades sin filtros implícitos**: cuando el '
        'user diga "dame las oportunidades de X mes", llama '
        '`oportunidades_por_periodo` SIN parámetros adicionales — trae '
        'TODAS, sin filtrar por etapa ni monto. Solo filtra si el user lo '
        'pide explícitamente ("las ganadas de mayo", "las mayores a $1M").\n'
        '6e. **"Más prometedora"**: si el user dice "la más prometedora" '
        'sin especificar periodo, usa `top_oportunidades_prometedoras` con '
        'defaults (filtra por mes_cierre = mes actual). Si dice "del mes '
        'pasado" o "de marzo", pasa mes/anio. Si dice "de todo el '
        'pipeline" pasa todos_los_meses=true. En la respuesta SIEMPRE di '
        'a qué periodo corresponde ("de las que cierran en mayo").\n'
        '7. **NO uses emojis decorativos** (sin "✨", "🎉", "💰", etc.). '
        'Si quieres énfasis usa **negritas**. Los símbolos como ↗ o ✓ están '
        'bien si suman información.\n'
        '8. **No te disculpes en exceso.** Corrige y sigue.\n\n'

        '## Cómo respondes\n'
        '- **Pregunta corta** (saludo, "qué tal"): respuesta corta, sin menús.\n'
        '- **Pregunta de un dato**: respuesta en 1-3 líneas con el número '
        'destacado en **negritas**.\n'
        '- **Pregunta tipo reporte** ("cómo va el mes", "ranking"): heading '
        '`### Título` + bullets o tabla con los KPIs clave + 1-2 líneas de '
        'lectura analítica al final ("destaca X", "ojo con Y"). Como un '
        'consultor entregando un brief.\n'
        '- **Lista de oportunidades**: bullets con cada opp como link '
        '`[Título — Cliente — $Monto](opp:ID)`. Agrupa por vendedor o etapa '
        'si tiene sentido.\n'
        '- **Evaluación de un vendedor**: usa `detalle_vendedor` y arma una '
        'mini-evaluación: pipeline, monto ganado del mes vs año, # clientes, '
        'lectura ("tiene buen ritmo", "está debajo del promedio del equipo").\n\n'

        '## Personalidad (sutil)\n'
        '- Saluda por nombre en el primer mensaje del día.\n'
        '- Mexicanismos suaves OK ("va", "órale", "andamos") con moderación.\n'
        '- Números buenos: una línea de reconocimiento breve. Números malos: '
        'sé honesto pero constructivo ("este mes fue retador; el pipeline '
        'sigue sólido para junio").\n'
        '- Si te preguntan algo fuera del CRM, redirige con cortesía.\n'
    )
    return {'role': 'system', 'content': base + _user_context_block(user)}


# ─── Endpoints ────────────────────────────────────────────────────────


@login_required
@require_http_methods(['GET'])
def api_asistente_config(request):
    """Devuelve la config visible (nombre + logo) para que el frontend la pinte."""
    cfg = AsistenteConfig.get_singleton()
    logo_url = cfg.logo.url if cfg.logo else ''
    return JsonResponse({
        'ok': True,
        'nombre': cfg.nombre,
        'logo_url': logo_url,
        'activo': cfg.activo,
    })


@login_required
@require_http_methods(['GET'])
def api_asistente_conversacion(request):
    """Devuelve el historial COMPLETO de la conversación del user (para mostrar en UI)."""
    conv = _get_or_create_conv(request.user)
    # Filtramos: solo mostramos al user los mensajes user + assistant final.
    # Los tool/assistant-con-tool-call quedan ocultos (son ruido para el user).
    msgs = []
    for m in conv.mensajes.all():
        if m.role == 'user':
            msgs.append({
                'id': m.id,
                'role': 'user',
                'contenido': m.contenido,
                'fecha': m.fecha.isoformat(),
            })
        elif m.role == 'assistant' and not m.tool_name and m.contenido:
            msgs.append({
                'id': m.id,
                'role': 'assistant',
                'contenido': m.contenido,
                'fecha': m.fecha.isoformat(),
            })
    return JsonResponse({'ok': True, 'mensajes': msgs})


@login_required
@require_http_methods(['DELETE'])
def api_asistente_conversacion_eliminar(request):
    """Limpia el historial completo del user. Empieza la conversación de cero."""
    conv = _get_or_create_conv(request.user)
    conv.mensajes.all().delete()
    return JsonResponse({'ok': True})


@login_required
@require_http_methods(['GET'])
def api_asistente_reporte_xlsx(request):
    """Genera un reporte Excel a partir de un token de un solo uso creado
    por la tool `generar_reporte_excel`. El token vive 1 hora y guarda los
    filtros + el user_id que lo solicitó (para preservar permisos).
    """
    from django.core.cache import cache
    from django.http import HttpResponse
    token = (request.GET.get('token') or '').strip()
    if not token:
        return JsonResponse({'ok': False, 'error': 'Falta token.'}, status=400)
    payload = cache.get('asist_xlsx_' + token)
    if not payload:
        return JsonResponse({'ok': False, 'error': 'Token expirado o inválido.'}, status=410)
    if payload.get('user_id') != request.user.id:
        return JsonResponse({'ok': False, 'error': 'Token no pertenece a este usuario.'}, status=403)

    tipo = payload.get('tipo') or 'oportunidades'
    if tipo != 'oportunidades':
        return JsonResponse({'ok': False, 'error': f'Tipo de reporte no soportado: {tipo}'}, status=400)

    # Construir queryset con permisos del user solicitante.
    from .models import TodoItem
    from .views_utils import is_supervisor
    from .views_grupos import get_usuarios_visibles_ids
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    mes = payload.get('mes')
    anio = payload.get('anio')
    vendedor_username = payload.get('vendedor_username')

    qs = TodoItem.objects.select_related('cliente', 'usuario')
    if mes and anio:
        qs = qs.filter(fecha_creacion__year=anio, fecha_creacion__month=mes)
    if not (is_supervisor(request.user) or request.user.is_superuser):
        try:
            ids = list(get_usuarios_visibles_ids(request.user)) or [request.user.id]
            qs = qs.filter(usuario_id__in=ids)
        except Exception:
            qs = qs.filter(usuario=request.user)
    if vendedor_username:
        from django.contrib.auth.models import User as _User
        try:
            v = _User.objects.get(username=vendedor_username)
            qs = qs.filter(usuario=v)
        except _User.DoesNotExist:
            pass

    qs = qs.order_by('-monto')

    # Construir el workbook
    wb = Workbook()
    ws = wb.active
    ws.title = f'Oportunidades {mes or "X"}-{anio or "X"}'

    headers = ['ID', 'Oportunidad', 'Cliente', 'Vendedor', 'Monto MXN',
               'Probabilidad %', 'Etapa', 'Producto', 'Área',
               'Tipo negociación', 'Mes cierre', 'Año cierre',
               'Fecha creación', 'PO', 'Factura']
    header_fill = PatternFill('solid', fgColor='0066FF')
    header_font = Font(bold=True, color='FFFFFF', size=11)
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row_idx, opp in enumerate(qs, start=2):
        ws.cell(row=row_idx, column=1, value=opp.id)
        ws.cell(row=row_idx, column=2, value=opp.oportunidad)
        ws.cell(row=row_idx, column=3, value=opp.cliente.nombre_empresa if opp.cliente_id else '')
        ws.cell(row=row_idx, column=4, value=(opp.usuario.get_full_name() or opp.usuario.username) if opp.usuario_id else '')
        ws.cell(row=row_idx, column=5, value=float(opp.monto or 0))
        ws.cell(row=row_idx, column=6, value=opp.probabilidad_cierre or 0)
        ws.cell(row=row_idx, column=7, value=opp.etapa_corta or opp.etapa_completa or '')
        ws.cell(row=row_idx, column=8, value=opp.producto or '')
        ws.cell(row=row_idx, column=9, value=opp.area or '')
        ws.cell(row=row_idx, column=10, value=opp.tipo_negociacion or '')
        ws.cell(row=row_idx, column=11, value=opp.mes_cierre or '')
        ws.cell(row=row_idx, column=12, value=opp.anio_cierre or '')
        ws.cell(row=row_idx, column=13, value=opp.fecha_creacion.strftime('%Y-%m-%d') if opp.fecha_creacion else '')
        ws.cell(row=row_idx, column=14, value=opp.po_number or '')
        ws.cell(row=row_idx, column=15, value=opp.factura_numero or '')

    # Anchos automáticos rough
    widths = [8, 36, 28, 24, 14, 8, 18, 14, 14, 14, 8, 8, 12, 14, 14]
    for i, w in enumerate(widths, start=1):
        col_letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[col_letter].width = w

    # Formato moneda en col E (Monto)
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=5).number_format = '"$"#,##0.00'

    ws.freeze_panes = 'A2'

    # Servir como descarga
    import io
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f'oportunidades_{mes or "all"}-{anio or "all"}.xlsx'
    resp = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    resp['Content-Disposition'] = f'attachment; filename="{filename}"'
    # Token de un solo uso: invalidamos después de la descarga.
    cache.delete('asist_xlsx_' + token)
    return resp


@login_required
@require_http_methods(['POST'])
def api_asistente_mensaje(request):
    """Recibe mensaje del user, llama al LLM (con tool use loop), devuelve respuesta.

    Body: { "texto": "string" }
    """
    try:
        data = json.loads(request.body or b'{}')
    except json.JSONDecodeError:
        return JsonResponse({'ok': False, 'error': 'JSON inválido'}, status=400)

    texto = (data.get('texto') or '').strip()
    if not texto:
        return JsonResponse({'ok': False, 'error': 'Mensaje vacío'}, status=400)
    if len(texto) > 4000:
        return JsonResponse({'ok': False, 'error': 'Mensaje demasiado largo (>4000 chars)'}, status=400)

    cfg = AsistenteConfig.get_singleton()
    if not cfg.activo:
        return JsonResponse({'ok': False, 'error': 'El asistente está deshabilitado.'}, status=503)

    conv = _get_or_create_conv(request.user)

    # 1) Guardar mensaje del user
    user_msg = MensajeAsistente.objects.create(
        conversacion=conv, role='user', contenido=texto,
    )

    # 2) Loop de tool use
    tools = tools_for_user(request.user)
    sys_msg = _system_prompt(request.user, cfg)
    final_text = None

    try:
        for iteration in range(MAX_TOOL_ITERATIONS):
            # Reconstruir contexto en cada iteración (incluye los tool_results
            # recién agregados).
            context_msgs = _build_context(conv, cfg)
            messages = [sys_msg] + context_msgs

            resp = chat(
                messages=messages,
                tools=tools,
                model=cfg.modelo,
                temperature=0.3,
                max_tokens=1024,
            )

            # Si el modelo regresó texto sin pedir tools, terminamos.
            if not resp.get('tool_calls'):
                final_text = (resp.get('text') or '').strip() or '(sin respuesta)'
                MensajeAsistente.objects.create(
                    conversacion=conv, role='assistant', contenido=final_text,
                )
                break

            # El modelo pidió tools. Guardar mensaje assistant con tool_calls,
            # luego ejecutar cada tool y guardar sus resultados.
            for tc in resp['tool_calls']:
                MensajeAsistente.objects.create(
                    conversacion=conv,
                    role='assistant',
                    contenido=resp.get('text') or '',
                    tool_name=tc['name'],
                    tool_args_json=tc.get('arguments_str') or json.dumps(tc.get('arguments') or {}),
                    tool_call_id=tc['id'],
                )
                tool_result = execute_tool(tc['name'], tc.get('arguments') or {}, request.user)
                MensajeAsistente.objects.create(
                    conversacion=conv,
                    role='tool',
                    contenido=json.dumps(tool_result, ensure_ascii=False, default=str),
                    tool_name=tc['name'],
                    tool_call_id=tc['id'],
                )
            # Siguiente iteración: el modelo verá los resultados y responderá.
        else:
            # Salió del for sin break → llegó al tope sin texto final.
            final_text = (
                'Lo siento, no pude llegar a una respuesta después de varias '
                'consultas. Intenta reformular tu pregunta.'
            )
            MensajeAsistente.objects.create(
                conversacion=conv, role='assistant', contenido=final_text,
            )

    except AsistenteError as e:
        # Borrar el user_msg para que el user pueda reintentar sin duplicar.
        # No: lo dejamos en BD para que vea su intento. Solo devolvemos error.
        log.warning('Asistente error: %s', e)
        return JsonResponse({'ok': False, 'error': str(e)}, status=502)
    except Exception as e:
        log.exception('Error inesperado en asistente: %s', e)
        return JsonResponse({'ok': False, 'error': f'Error inesperado: {e}'}, status=500)

    return JsonResponse({
        'ok': True,
        'mensaje_user_id': user_msg.id,
        'respuesta': final_text or '(sin respuesta)',
    })
